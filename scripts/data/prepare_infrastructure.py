from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZipFile
import io

import pandas as pd
import shapefile
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw"
OUTDIR = ROOT / "data" / "interim"
OUTPUT = OUTDIR / "mart_infraestructura_municipio_veracruz.csv"

DENUE_72_FILES = [RAW / "denue_00_72_1_csv.zip", RAW / "denue_00_72_2_csv.zip"]
DENUE_71_FILE = RAW / "denue_00_71_csv.zip"
DENUE_48_49_FILE = RAW / "denue_00_48-49_csv.zip"
DENUE_56_FILE = RAW / "denue_00_56_shp.zip"
PIB_FILE = RAW / "PIB_Turistico_Estatal_y_Municipal_2018-2022_.zip"

VERACRUZ_CODE = "30"
VERACRUZ_NAME = "Veracruz de Ignacio de la Llave"

# Explicit SCIAN code sets used by the original validated municipal mart.
# Keeping them visible makes the proxy definitions inspectable and testable.
RECREATION_TOURISM_PROXY_CODES = {
    "711312",  # public promoters with facilities
    "712111",  # private museums
    "712112",  # public museums
    "712120",  # historical sites
    "712131",  # private botanical gardens / zoos
    "712132",  # public botanical gardens / zoos
    "712190",  # caves, natural parks and cultural-heritage sites
    "713113",  # private water parks / spas
    "713114",  # public water parks / spas
    "713998",  # other private recreational services
    "713999",  # other public recreational services
}

AGENCIES_TOURS_EVENTS_CODES = {
    "561510",  # travel agencies
    "561520",  # tour/package organization for travel agencies
    "561590",  # other reservation services
    "561920",  # convention/trade-fair organizers
}

TOURIST_TRANSPORT_CODES = {
    "485510",  # bus rental with driver
    "487110",  # tourist transport by land
    "487210",  # tourist transport by water
}


def _metadata_text(path: Path) -> str:
    with ZipFile(path) as zf:
        matches = [n for n in zf.namelist() if n.endswith("metadatos_denue.txt")]
        if not matches:
            raise ValueError(f"DENUE archive lacks metadata: {path.name}")
        return zf.read(matches[0]).decode("utf-8-sig", errors="replace")


def validate_denue_snapshot(paths: list[Path]) -> None:
    for path in paths:
        text = _metadata_text(path)
        if "DENUE-2026" not in text and "05_2026" not in text:
            raise ValueError(f"Expected DENUE 05/2026 snapshot in {path.name}")


def _add_municipal_key(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["cve_ent"] = out["cve_ent"].astype(str).str.strip().str.zfill(2)
    out["cve_mun"] = out["cve_mun"].astype(str).str.strip().str.zfill(3)
    out["clave_municipio"] = out["cve_ent"] + out["cve_mun"]
    out["municipio"] = out["municipio"].astype(str).str.strip()
    out["codigo_act"] = out["codigo_act"].astype(str).str.strip()
    return out


def _read_denue_csv_zip(path: Path) -> pd.DataFrame:
    with ZipFile(path) as zf:
        names = [
            n
            for n in zf.namelist()
            if "conjunto_de_datos/" in n and n.lower().endswith(".csv")
        ]
        if len(names) != 1:
            raise ValueError(f"Expected exactly one DENUE data CSV in {path.name}")
        with zf.open(names[0]) as fh:
            df = pd.read_csv(
                fh,
                encoding="latin1",
                dtype=str,
                usecols=["codigo_act", "nombre_act", "cve_ent", "cve_mun", "municipio"],
            )
    df = _add_municipal_key(df)
    return df[df["cve_ent"].eq(VERACRUZ_CODE)].copy()


def _read_denue_56_dbf(path: Path) -> pd.DataFrame:
    with ZipFile(path) as zf:
        names = [n for n in zf.namelist() if n.lower().endswith(".dbf")]
        if len(names) != 1:
            raise ValueError(f"Expected exactly one DENUE DBF in {path.name}")
        dbf_bytes = zf.read(names[0])

    with NamedTemporaryFile(suffix=".dbf") as tmp:
        tmp.write(dbf_bytes)
        tmp.flush()
        reader = shapefile.Reader(dbf=tmp.name, encoding="latin1")
        fields = [field[0] for field in reader.fields[1:]]
        positions = {name: i for i, name in enumerate(fields)}
        required = {"codigo_act", "nombre_act", "cve_ent", "cve_mun", "municipio"}
        missing = required - set(positions)
        if missing:
            raise ValueError(f"DENUE 56 DBF missing fields: {sorted(missing)}")

        rows = []
        for record in reader.iterRecords():
            cve_ent = str(record[positions["cve_ent"]]).strip().zfill(2)
            if cve_ent != VERACRUZ_CODE:
                continue
            rows.append(
                {
                    "codigo_act": str(record[positions["codigo_act"]]).strip(),
                    "nombre_act": str(record[positions["nombre_act"]]).strip(),
                    "cve_ent": cve_ent,
                    "cve_mun": str(record[positions["cve_mun"]]).strip().zfill(3),
                    "municipio": str(record[positions["municipio"]]).strip(),
                }
            )
    return _add_municipal_key(pd.DataFrame(rows))


def _count_by_key(df: pd.DataFrame, mask: pd.Series, name: str) -> pd.Series:
    return df.loc[mask].groupby("clave_municipio").size().rename(name)


def _load_pib_2022(path: Path) -> pd.DataFrame:
    workbook_name = "PIB_Turistico_Estatal_y_Municipal_2018-2022.xlsx"
    with ZipFile(path) as zf:
        if workbook_name not in zf.namelist():
            raise ValueError(f"{path.name} does not contain {workbook_name}")
        workbook_bytes = zf.read(workbook_name)

    wb = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    if "PIB_Municipal" not in wb.sheetnames:
        raise ValueError("PIB workbook does not contain the PIB_Municipal sheet")
    ws = wb["PIB_Municipal"]

    rows = []
    for values in ws.iter_rows(min_row=9, values_only=True):
        if values[3] != VERACRUZ_NAME:
            continue
        rows.append(
            {
                "clave_municipio": str(values[4]).strip().zfill(5),
                "municipio_pib": str(values[5]).strip(),
                "pib_municipal_2022": values[29],
                "pib_turistico_2022": values[31],
                "participacion_turismo_pct": values[33],
            }
        )
    out = pd.DataFrame(rows)
    if len(out) != 212 or not out["clave_municipio"].is_unique:
        raise AssertionError(
            f"Expected 212 unique Veracruz municipalities in PIB source, got {len(out)}"
        )
    return out


def _canonical_denue_names(*frames: pd.DataFrame) -> pd.DataFrame:
    names = pd.concat(
        [df[["clave_municipio", "municipio"]] for df in frames if not df.empty],
        ignore_index=True,
    ).drop_duplicates()
    counts = names.groupby("clave_municipio")["municipio"].nunique()
    ambiguous = counts[counts > 1]
    if not ambiguous.empty:
        details = names[names["clave_municipio"].isin(ambiguous.index)]
        raise AssertionError(f"DENUE municipality code maps to multiple names:\n{details}")
    return names.drop_duplicates("clave_municipio")


def build_infrastructure() -> pd.DataFrame:
    validate_denue_snapshot(
        DENUE_72_FILES + [DENUE_71_FILE, DENUE_48_49_FILE, DENUE_56_FILE]
    )

    d72 = pd.concat([_read_denue_csv_zip(path) for path in DENUE_72_FILES], ignore_index=True)
    d71 = _read_denue_csv_zip(DENUE_71_FILE)
    d48 = _read_denue_csv_zip(DENUE_48_49_FILE)
    d56 = _read_denue_56_dbf(DENUE_56_FILE)
    pib = _load_pib_2022(PIB_FILE)
    names = _canonical_denue_names(d72, d71, d48, d56)

    alojamiento = _count_by_key(
        d72, d72["codigo_act"].str.startswith("721", na=False), "alojamiento_n"
    )
    alimentos = _count_by_key(
        d72, d72["codigo_act"].str.startswith("722", na=False), "alimentos_bebidas_n"
    )
    recreacion_total = d71.groupby("clave_municipio").size().rename("sector_recreacion_total_n")
    recreacion_proxy = _count_by_key(
        d71,
        d71["codigo_act"].isin(RECREATION_TOURISM_PROXY_CODES),
        "recreacion_turistica_proxy_n",
    )
    agencias = _count_by_key(
        d56,
        d56["codigo_act"].isin(AGENCIES_TOURS_EVENTS_CODES),
        "agencias_tours_eventos_n",
    )
    transporte = _count_by_key(
        d48,
        d48["codigo_act"].isin(TOURIST_TRANSPORT_CODES),
        "transporte_turistico_n",
    )

    out = pib.set_index("clave_municipio")
    for series in [alojamiento, alimentos, recreacion_total, recreacion_proxy, agencias, transporte]:
        out = out.join(series, how="left")
    out = out.reset_index().merge(names, on="clave_municipio", how="left", validate="one_to_one")
    out["municipio"] = out["municipio"].fillna(out["municipio_pib"])

    count_cols = [
        "alojamiento_n",
        "alimentos_bebidas_n",
        "sector_recreacion_total_n",
        "recreacion_turistica_proxy_n",
        "agencias_tours_eventos_n",
        "transporte_turistico_n",
    ]
    out[count_cols] = out[count_cols].fillna(0).astype("int64")

    if (out[count_cols] < 0).any().any():
        raise AssertionError("DENUE counts must be non-negative")
    if (out["recreacion_turistica_proxy_n"] > out["sector_recreacion_total_n"]).any():
        raise AssertionError("Recreation proxy cannot exceed the full sector-71 count")

    columns = [
        "clave_municipio",
        "municipio",
        *count_cols,
        "pib_turistico_2022",
        "participacion_turismo_pct",
        "pib_municipal_2022",
    ]
    return out[columns].sort_values("clave_municipio", kind="stable").reset_index(drop=True)


def main() -> None:
    OUTDIR.mkdir(parents=True, exist_ok=True)
    out = build_infrastructure()
    out.to_csv(OUTPUT, index=False)
    print(f"Wrote {len(out):,} rows -> {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
